"""
Leave-one-well-out cross-validation for facies classification.

Each fold: one well as test, the remaining wells as training.
Results (accuracy + per-depth prediction CSV) are saved under --base_output_dir.

Example usage:
    python run_cross_val.py \
        --xlsx_path ./facies-sand-gr-diff.xlsx \
        --model convnext_tiny \
        --window_size 100 \
        --window_stride 0 \
        --input_size 224 \
        --batch_size 64 \
        --epochs 100 \
        --lr 1e-3 \
        --base_output_dir ./outputs/cross_val
"""

import argparse
import csv
import json
import os
import subprocess
import sys
from pathlib import Path

from welllog.well_split import get_well_classes, check_train_covers_all_classes

# ── well list (auto-detected from xlsx if not overridden) ─────────────────────
def detect_wells(xlsx_path):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True)
    return [s for s in wb.sheetnames if s.lower() != "mapping"]


def get_args_parser():
    p = argparse.ArgumentParser("Leave-one-well-out cross-validation")

    # Data
    # p.add_argument("--xlsx_path", default="./facies-sand-gr-diff.xlsx")
    p.add_argument("--xlsx_path", default="./facies-gr-diff.xlsx")
    p.add_argument("--feature_cols", default="GR,GR_diff1,GR_diff2")
    p.add_argument("--label_col", default="facies")
    p.add_argument("--depth_col", default="DEPT")
    p.add_argument("--wells", default="", help="comma-separated well list; empty = auto-detect from xlsx")

    # Window
    p.add_argument("--window_size", default=100, type=int)
    p.add_argument("--window_stride", default=0, type=int, help="0 = auto (window_size//2)")
    p.add_argument("--input_size", default=224, type=int)

    # Model
    p.add_argument("--model", default="convnext_tiny")
    p.add_argument("--drop_path", default=0.0, type=float)

    # Training
    p.add_argument("--batch_size", default=64, type=int)
    p.add_argument("--epochs", default=100, type=int)
    p.add_argument("--lr", default=1e-3, type=float)
    p.add_argument("--weight_decay", default=0.05, type=float)
    p.add_argument("--warmup_epochs", default=20, type=int)
    p.add_argument("--device", default="cuda")
    p.add_argument("--seed", default=42, type=int)
    p.add_argument("--num_workers", default=4, type=int)
    p.add_argument("--use_amp", default=False, action="store_true")

    # Output / post-process (forwarded to main.py)
    p.add_argument("--min_segment_length", default=5, type=int,
                   help="short-segment merge for predictions.csv; 1 = off")
    p.add_argument("--base_output_dir", default="./outputs/cross_val")
    p.add_argument("--python", default=sys.executable, help="python executable to use")

    return p


def run_fold(fold_idx, test_well, train_wells, args):
    """Run one fold via subprocess calling main.py."""
    fold_name = f"fold_{fold_idx:02d}_{test_well}"
    output_dir = os.path.join(args.base_output_dir, fold_name)
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    train_str = ",".join(train_wells)
    cmd = [
        args.python, "main.py",
        "--data_set", "WELLLOG_XLSX",
        "--xlsx_path", args.xlsx_path,
        "--feature_cols", args.feature_cols,
        "--label_col", args.label_col,
        "--depth_col", args.depth_col,
        "--train_wells", train_str,
        "--val_wells", test_well,
        "--window_size", str(args.window_size),
        "--window_stride", str(args.window_stride),
        "--input_size", str(args.input_size),
        "--model", args.model,
        "--drop_path", str(args.drop_path),
        "--batch_size", str(args.batch_size),
        "--epochs", str(args.epochs),
        "--lr", str(args.lr),
        "--weight_decay", str(args.weight_decay),
        "--warmup_epochs", str(args.warmup_epochs),
        "--device", args.device,
        "--seed", str(args.seed),
        "--num_workers", str(args.num_workers),
        "--use_amp", "true" if args.use_amp else "false",
        "--mixup", "0",
        "--cutmix", "0",
        "--smoothing", "0",
        "--min_segment_length", str(args.min_segment_length),
        "--output_dir", output_dir,
        "--save_ckpt", "true",
        "--auto_resume", "false",
        "--dist_eval", "false",
    ]

    log_file = os.path.join(output_dir, "train.log")
    print(f"\n{'='*60}")
    print(f"[Fold {fold_idx}] Test well: {test_well}")
    print(f"           Train wells: {train_str}")
    print(f"           Output: {output_dir}")
    print(f"{'='*60}")

    with open(log_file, "w") as lf:
        proc = subprocess.run(cmd, stdout=lf, stderr=subprocess.STDOUT, cwd=os.path.dirname(os.path.abspath(__file__)))

    # ── parse best accuracy from log.txt ─────────────────────────────────────
    best_acc = None
    main_log = os.path.join(output_dir, "log.txt")
    if os.path.exists(main_log):
        with open(main_log) as f:
            for line in f:
                try:
                    entry = json.loads(line)
                    acc = entry.get("test_acc1")
                    if acc is not None:
                        if best_acc is None or acc > best_acc:
                            best_acc = acc
                except json.JSONDecodeError:
                    pass

    # ── check prediction csv ──────────────────────────────────────────────────
    pred_csv = os.path.join(output_dir, "predictions.csv")
    pred_rows = 0
    if os.path.exists(pred_csv):
        with open(pred_csv) as f:
            pred_rows = sum(1 for _ in f) - 1  # exclude header

    result = {
        "fold": fold_idx,
        "test_well": test_well,
        "train_wells": train_str,
        "best_acc1": best_acc,
        "pred_csv": pred_csv if os.path.exists(pred_csv) else None,
        "pred_rows": pred_rows,
        "returncode": proc.returncode,
    }
    print(f"[Fold {fold_idx}] Done  best_acc1={best_acc}  pred_rows={pred_rows}  rc={proc.returncode}")
    return result


def main():
    parser = get_args_parser()
    args = parser.parse_args()

    # ── resolve well list ─────────────────────────────────────────────────────
    if args.wells:
        all_wells = [w.strip() for w in args.wells.split(",") if w.strip()]
    else:
        all_wells = detect_wells(args.xlsx_path)
    print(f"Wells detected: {all_wells}  ({len(all_wells)} folds)")

    Path(args.base_output_dir).mkdir(parents=True, exist_ok=True)

    # ── pre-load class sets for coverage check ────────────────────────────────
    well_classes = get_well_classes(args.xlsx_path, args.label_col)
    all_classes = set().union(*well_classes.values()) if well_classes else set()

    # ── run each fold ─────────────────────────────────────────────────────────
    results = []
    for i, test_well in enumerate(all_wells):
        train_wells = [w for w in all_wells if w != test_well]
        if not check_train_covers_all_classes(train_wells, well_classes, all_classes):
            print(
                f"\n[Fold {i}] SKIP: test well '{test_well}' is the sole source of "
                f"at least one class — train wells would not cover all classes."
            )
            results.append({
                "fold": i,
                "test_well": test_well,
                "train_wells": ",".join(train_wells),
                "best_acc1": None,
                "pred_csv": None,
                "pred_rows": 0,
                "returncode": -1,
            })
            continue
        result = run_fold(i, test_well, train_wells, args)
        results.append(result)

    # ── print summary ─────────────────────────────────────────────────────────
    print(f"\n{'='*60}")
    print("Cross-validation summary")
    print(f"{'='*60}")
    header = f"{'Fold':>4}  {'Test well':<14}  {'Acc@1':>7}  {'Pred rows':>9}  {'RC':>3}"
    print(header)
    print("-" * len(header))
    accs = []
    for r in results:
        acc_str = f"{r['best_acc1']:.2f}%" if r["best_acc1"] is not None else "   N/A"
        print(f"{r['fold']:>4}  {r['test_well']:<14}  {acc_str:>7}  {r['pred_rows']:>9}  {r['returncode']:>3}")
        if r["best_acc1"] is not None:
            accs.append(r["best_acc1"])
    print("-" * len(header))
    if accs:
        mean_acc = sum(accs) / len(accs)
        print(f"       Mean Acc@1: {mean_acc:.2f}%  ({len(accs)}/{len(results)} folds succeeded)")

    # ── save summary csv ──────────────────────────────────────────────────────
    summary_path = os.path.join(args.base_output_dir, "cross_val_summary.csv")
    with open(summary_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["fold", "test_well", "train_wells",
                                               "best_acc1", "pred_rows", "pred_csv", "returncode"])
        writer.writeheader()
        writer.writerows(results)
    print(f"\nSummary saved -> {summary_path}")

    # ── merge all prediction CSVs into one ───────────────────────────────────
    merged_path = os.path.join(args.base_output_dir, "all_predictions.csv")
    wrote_header = False
    with open(merged_path, "w", newline="", encoding="utf-8") as out:
        for r in results:
            if r["pred_csv"] and os.path.exists(r["pred_csv"]):
                with open(r["pred_csv"], newline="", encoding="utf-8") as inp:
                    reader = csv.DictReader(inp)
                    if not wrote_header:
                        writer = csv.DictWriter(out, fieldnames=reader.fieldnames)
                        writer.writeheader()
                        wrote_header = True
                    for row in reader:
                        writer.writerow(row)
    if wrote_header:
        print(f"All predictions merged -> {merged_path}")


if __name__ == "__main__":
    main()
