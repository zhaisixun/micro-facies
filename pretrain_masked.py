#!/usr/bin/env python3
"""Masked self-supervised pretraining and full-well reconstruction export."""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook
from torch.utils.data import DataLoader

from models.masked_welllog import MaskedWellLogAutoencoder
from welllog.masked_pretrain import MaskedWellLogWindowDataset, load_unlabeled_wells, window_starts


def parse_args():
    parser = argparse.ArgumentParser("Masked well-log reconstruction pretraining")
    parser.add_argument("--xlsx_path", default="./facies-gr-diff0614-用GR-CNL-DEN.xlsx")
    parser.add_argument("--feature_cols", default="GR,CNL,DEN")
    parser.add_argument("--depth_col", default="DEPT")
    parser.add_argument("--train_wells", default="",
                        help="Optional manual training wells. Leave empty for automatic well-level split.")
    parser.add_argument("--val_wells", default="",
                        help="Optional manual validation wells. Leave empty for automatic well-level split.")
    parser.add_argument("--val_ratio", type=float, default=0.2,
                        help="Validation-well fraction used by automatic split.")
    parser.add_argument("--window_size", type=int, default=128)
    parser.add_argument("--window_stride", type=int, default=32)
    parser.add_argument("--channel_mask_prob", type=float, default=0.25)
    parser.add_argument("--depth_mask_prob", type=float, default=0.75)
    parser.add_argument("--depth_span", type=int, default=16)
    parser.add_argument("--depth_spans", type=int, default=2)
    parser.add_argument("--encoder_output_stride", type=int, choices=[8, 16, 32], default=8)
    parser.add_argument("--decoder_channels", type=int, default=256)
    parser.add_argument("--batch_size", type=int, default=64)
    parser.add_argument("--epochs", type=int, default=100)
    parser.add_argument("--lr", type=float, default=1e-3)
    parser.add_argument("--weight_decay", type=float, default=0.05)
    parser.add_argument("--num_workers", type=int, default=4)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--device", default="cuda")
    parser.add_argument("--output_dir", default="./outputs/masked_pretrain")
    parser.add_argument("--resume", default="")
    parser.add_argument("--eval_only", action="store_true")
    parser.add_argument("--mape_epsilon", type=float, default=1e-6)
    return parser.parse_args()


def split_names(value):
    return [item.strip() for item in value.split(",") if item.strip()]


def resolve_well_split(args):
    """Resolve an explicit split or create a deterministic split by workbook sheet."""
    manual_train = split_names(args.train_wells)
    manual_val = split_names(args.val_wells)
    if bool(manual_train) != bool(manual_val):
        raise ValueError("Set both --train_wells and --val_wells, or leave both empty")
    if manual_train:
        overlap = sorted(set(manual_train) & set(manual_val))
        if overlap:
            raise ValueError(f"Train/validation wells overlap: {overlap}")
        return manual_train, manual_val

    if not 0.0 < args.val_ratio < 1.0:
        raise ValueError("--val_ratio must be between 0 and 1")
    workbook = load_workbook(args.xlsx_path, read_only=True, data_only=True)
    wells = list(workbook.sheetnames)
    workbook.close()
    if len(wells) < 2:
        raise ValueError("Automatic split requires at least two wells")
    rng = random.Random(args.seed)
    rng.shuffle(wells)
    val_count = min(max(1, round(len(wells) * args.val_ratio)), len(wells) - 1)
    val_wells = sorted(wells[:val_count])
    train_wells = sorted(wells[val_count:])
    return train_wells, val_wells


def run_epoch(model, loader, device, optimizer=None):
    training = optimizer is not None
    model.train(training)
    loss_sum, masked_count = 0.0, 0
    for target, visible in loader:
        target = target.to(device, non_blocking=True)
        visible = visible.to(device, non_blocking=True)
        with torch.set_grad_enabled(training):
            prediction = model(target, visible)
            loss = model.masked_loss(prediction, target, visible)
        if training:
            optimizer.zero_grad(set_to_none=True)
            loss.backward()
            torch.nn.utils.clip_grad_norm_(model.parameters(), 5.0)
            optimizer.step()
        count = int((~visible).sum().item())  # 
        loss_sum += float(loss.item()) * count
        masked_count += count
    return loss_sum / max(masked_count, 1)


@torch.no_grad()
def reconstruct_well(model, store, window_size, stride, depth_span, device):
    """Predict every point while hidden, then fuse overlapping windows."""
    feat = store["feat"]
    channels = feat.shape[1]
    prediction_sum = np.zeros_like(feat, dtype=np.float64)
    weight_sum = np.zeros_like(feat, dtype=np.float64)
    local_weight = np.maximum(np.hanning(window_size).astype(np.float32), 0.05)
    model.eval()
    for start in window_starts(len(feat), window_size, stride):
        window = torch.from_numpy(feat[start:start + window_size].T.copy()).float()
        inputs, masks, regions = [], [], []
        for channel in range(channels):
            for span_start in range(0, window_size, depth_span):
                span_end = min(span_start + depth_span, window_size)
                visible = torch.ones_like(window, dtype=torch.bool)
                visible[channel, span_start:span_end] = False
                inputs.append(window)
                masks.append(visible)
                regions.append((channel, span_start, span_end))
        predictions = model(
            torch.stack(inputs).to(device), torch.stack(masks).to(device)
        ).cpu().numpy()
        for row, (channel, span_start, span_end) in enumerate(regions):
            destination = slice(start + span_start, start + span_end)
            weights = local_weight[span_start:span_end]
            prediction_sum[destination, channel] += predictions[row, channel, span_start:span_end] * weights
            weight_sum[destination, channel] += weights
    if np.any(weight_sum == 0):
        raise RuntimeError("Reconstruction left uncovered depth points")
    return prediction_sum / weight_sum


def metric_rows(original, prediction, curve_names, epsilon, well):
    rows = []
    for curve_idx, curve in enumerate(curve_names):
        true = original[:, curve_idx].astype(np.float64)
        pred = prediction[:, curve_idx].astype(np.float64)
        error = pred - true
        rows.append({
            "well": well, "curve": curve, "count": len(true),
            "mae": float(np.mean(np.abs(error))),
            "mse": float(np.mean(error ** 2)),
            "mape_percent": float(np.mean(np.abs(error) / np.maximum(np.abs(true), epsilon)) * 100.0),
        })
    return rows


def export_reconstructions(model, stores, args, curves, output_dir):
    rows, all_true, all_prediction = [], [], []
    for well, store in stores.items():
        prediction = reconstruct_well(
            model, store, args.window_size, args.window_stride, args.depth_span,
            next(model.parameters()).device,
        )
        original = store["feat"]
        all_true.append(original)
        all_prediction.append(prediction)
        rows.extend(metric_rows(original, prediction, curves, args.mape_epsilon, well))
        with (output_dir / f"reconstruction_{well}.csv").open("w", newline="", encoding="utf-8-sig") as handle:
            fields = ["well", "depth"]
            for curve in curves:
                fields.extend([f"original_{curve}", f"reconstructed_{curve}"])
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            for index, depth in enumerate(store["depths"]):
                row = {"well": well, "depth": depth}
                for curve_idx, curve in enumerate(curves):
                    row[f"original_{curve}"] = float(original[index, curve_idx])
                    row[f"reconstructed_{curve}"] = float(prediction[index, curve_idx])
                writer.writerow(row)
    rows.extend(metric_rows(
        np.concatenate(all_true), np.concatenate(all_prediction), curves,
        args.mape_epsilon, "ALL",
    ))
    with (output_dir / "reconstruction_metrics.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=["well", "curve", "count", "mae", "mse", "mape_percent"])
        writer.writeheader()
        writer.writerows(rows)
    return rows


def main():
    args = parse_args()
    random.seed(args.seed)
    np.random.seed(args.seed)
    torch.manual_seed(args.seed)
    if args.window_size <= 0 or args.window_stride <= 0 or args.depth_span <= 0:
        raise ValueError("window_size, window_stride and depth_span must be positive")
    device = torch.device(args.device if args.device != "cuda" or torch.cuda.is_available() else "cpu")
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    curves = split_names(args.feature_cols)
    train_wells, val_wells = resolve_well_split(args)
    print(f"Automatic/manual well split (seed={args.seed}):")
    print(f"  train ({len(train_wells)}): {train_wells}")
    print(f"  val   ({len(val_wells)}): {val_wells}")
    with (output_dir / "well_split.json").open("w", encoding="utf-8") as handle:
        json.dump(
            {"seed": args.seed, "val_ratio": args.val_ratio,
             "train_wells": train_wells, "val_wells": val_wells},
            handle, ensure_ascii=False, indent=2,
        )
    train_stores = load_unlabeled_wells(args.xlsx_path, train_wells, curves, args.depth_col)
    val_stores = load_unlabeled_wells(args.xlsx_path, val_wells, curves, args.depth_col)
    dataset_args = (
        args.window_size, args.window_stride, args.channel_mask_prob,
        args.depth_mask_prob, args.depth_span, args.depth_spans,
    )
    train_loader = DataLoader(
        MaskedWellLogWindowDataset(train_stores, *dataset_args),
        batch_size=args.batch_size, shuffle=True, num_workers=args.num_workers,
        pin_memory=device.type == "cuda",
    )
    val_loader = DataLoader(
        MaskedWellLogWindowDataset(val_stores, *dataset_args),
        batch_size=args.batch_size, shuffle=False, num_workers=args.num_workers,
        pin_memory=device.type == "cuda",
    )
    model = MaskedWellLogAutoencoder(
        in_chans=len(curves), decoder_channels=args.decoder_channels,
        encoder_output_stride=args.encoder_output_stride,
    ).to(device)
    optimizer = torch.optim.AdamW(model.parameters(), lr=args.lr, weight_decay=args.weight_decay)
    start_epoch, best_loss = 0, math.inf
    if args.resume:
        checkpoint = torch.load(args.resume, map_location="cpu")
        model.load_state_dict(checkpoint["model"])
        if "optimizer" in checkpoint and not args.eval_only:
            optimizer.load_state_dict(checkpoint["optimizer"])
        start_epoch = int(checkpoint.get("epoch", -1)) + 1
        best_loss = float(checkpoint.get("best_val_loss", math.inf))

    if not args.eval_only:
        for epoch in range(start_epoch, args.epochs):
            train_loss = run_epoch(model, train_loader, device, optimizer)
            random_state = random.getstate()
            random.seed(args.seed + 10000)
            val_loss = run_epoch(model, val_loader, device)
            random.setstate(random_state)
            record = {"epoch": epoch, "train_masked_smooth_l1": train_loss, "val_masked_smooth_l1": val_loss}
            print(json.dumps(record, ensure_ascii=False))
            with (output_dir / "log.jsonl").open("a", encoding="utf-8") as handle:
                handle.write(json.dumps(record, ensure_ascii=False) + "\n")
            is_best = val_loss < best_loss
            best_loss = min(best_loss, val_loss)
            checkpoint = {
                "model": model.state_dict(), "encoder": model.encoder_state_dict(),
                "optimizer": optimizer.state_dict(), "epoch": epoch,
                "best_val_loss": best_loss, "args": vars(args),
            }
            torch.save(checkpoint, output_dir / "checkpoint-last.pth")
            if is_best:
                torch.save(checkpoint, output_dir / "checkpoint-best.pth")

    best_path = Path(args.resume) if args.eval_only and args.resume else output_dir / "checkpoint-best.pth"
    if not best_path.exists():
        raise FileNotFoundError(f"No checkpoint available for reconstruction: {best_path}")
    model.load_state_dict(torch.load(best_path, map_location="cpu")["model"])
    model.to(device)
    for row in export_reconstructions(model, val_stores, args, curves, output_dir):
        if row["well"] == "ALL":
            print(json.dumps(row, ensure_ascii=False))


if __name__ == "__main__":
    main()
