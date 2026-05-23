"""
Utilities for automatically splitting wells into train / val sets.

The key constraint:
    Every class (label) that appears in the full dataset must appear in
    at least one training well.  Val wells may contain any subset of classes.

Algorithm
---------
1. Scan every well sheet and collect the set of classes each well contains.
2. Shuffle the well list with the given seed.
3. Greedily move wells into the val set as long as the remaining train wells
   still cover *all* classes.  Stop once the desired val_ratio is reached or
   no more wells can be safely moved.
4. Return (train_wells, val_wells).
"""

import math
import random
from typing import Dict, List, Set, Tuple


def _read_well_classes(xlsx_path: str, label_col: str) -> Dict[str, Set[str]]:
    """Return {well_name: set(class_strings)} for every non-mapping sheet."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    well_classes: Dict[str, Set[str]] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name.lower() == "mapping":
            continue
        ws = wb[sheet_name]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        if label_col not in header:
            continue
        label_idx = header.index(label_col)
        classes: Set[str] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            val = row[label_idx]
            if val is not None and str(val).strip() != "":
                classes.add(str(val).strip())
        if classes:
            well_classes[sheet_name] = classes
    wb.close()
    return well_classes


def auto_split_wells(
    xlsx_path: str,
    label_col: str,
    val_ratio: float = 0.2,
    seed: int = 42,
) -> Tuple[List[str], List[str]]:
    """Automatically split wells into (train_wells, val_wells).

    Parameters
    ----------
    xlsx_path  : path to the xlsx file (each sheet = one well)
    label_col  : column name that contains the class / facies label
    val_ratio  : target fraction of wells to put in val (0 < val_ratio < 1)
    seed       : random seed for reproducible shuffling

    Returns
    -------
    (train_wells, val_wells)  – both as lists of sheet-name strings

    Raises
    ------
    ValueError  if fewer than 2 wells are found, or if val_ratio is invalid
    """
    if not 0 < val_ratio < 1:
        raise ValueError(f"val_ratio must be in (0, 1), got {val_ratio}")

    well_classes = _read_well_classes(xlsx_path, label_col)
    if len(well_classes) < 2:
        raise ValueError(
            f"Need at least 2 wells in '{xlsx_path}' to perform auto split, "
            f"found {len(well_classes)}."
        )

    all_wells = list(well_classes.keys())
    rng = random.Random(seed)
    rng.shuffle(all_wells)

    all_classes: Set[str] = set().union(*well_classes.values())
    n_val_target = max(1, math.ceil(len(all_wells) * val_ratio))

    train_wells = list(all_wells)
    val_wells: List[str] = []

    for well in all_wells:
        if len(val_wells) >= n_val_target:
            break
        # Check if the remaining train wells (after removing this one) still
        # cover all classes.
        remaining_train = [w for w in train_wells if w != well]
        covered = set().union(*(well_classes[w] for w in remaining_train)) if remaining_train else set()
        if covered >= all_classes:
            train_wells.remove(well)
            val_wells.append(well)

    if not val_wells:
        raise ValueError(
            "Could not place any well in val while keeping full class coverage "
            "in train.  Consider a larger dataset or reducing val_ratio."
        )

    return train_wells, val_wells


def check_train_covers_all_classes(
    train_wells: List[str],
    well_classes: Dict[str, Set[str]],
    all_classes: Set[str],
) -> bool:
    """Return True if the given train wells cover every class in all_classes."""
    covered = set().union(*(well_classes[w] for w in train_wells if w in well_classes))
    return covered >= all_classes


def get_well_classes(xlsx_path: str, label_col: str) -> Dict[str, Set[str]]:
    """Public wrapper around _read_well_classes (for use in run_cross_val)."""
    return _read_well_classes(xlsx_path, label_col)
