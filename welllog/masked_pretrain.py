"""Datasets and masking utilities for self-supervised well-log reconstruction."""

from __future__ import annotations

import random
from dataclasses import dataclass
from typing import Dict, List

import numpy as np
import torch
from openpyxl import load_workbook
from torch.utils.data import Dataset


def load_unlabeled_wells(xlsx_path: str, wells: List[str], feature_cols: List[str], depth_col: str):
    """Load and normalize wells without requiring facies labels."""
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    stores: Dict[str, dict] = {}
    for well in wells:
        if well not in workbook.sheetnames:
            raise ValueError(f"Well '{well}' not found in {xlsx_path}")
        sheet = workbook[well]
        header = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        index = {name: i for i, name in enumerate(header)}
        missing = [name for name in feature_cols if name not in index]
        if missing:
            raise ValueError(f"Columns {missing} not found in well '{well}'")
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        raw = np.full((len(rows), len(feature_cols)), np.nan, dtype=np.float32)
        depths = []
        for row_idx, row in enumerate(rows):
            depths.append(row[index[depth_col]] if depth_col in index else row_idx)
            for col_idx, name in enumerate(feature_cols):
                value = row[index[name]]
                if value is not None and str(value).strip():
                    raw[row_idx, col_idx] = float(value)
        for col_idx in range(raw.shape[1]):
            median = np.nanmedian(raw[:, col_idx])
            raw[np.isnan(raw[:, col_idx]), col_idx] = 0.0 if np.isnan(median) else median
        mean = raw.mean(axis=0).astype(np.float32)
        std = raw.std(axis=0).astype(np.float32)
        std[std < 1e-6] = 1.0
        stores[well] = {
            "raw": raw,
            "feat": ((raw - mean) / std).astype(np.float32),
            "mean": mean,
            "std": std,
            "depths": depths,
        }
    return stores


@dataclass(frozen=True)
class WindowRef:
    well: str
    start: int


def window_starts(length: int, window_size: int, stride: int) -> List[int]:
    if length < window_size:
        return []
    starts = list(range(0, length - window_size + 1, stride))  # 每个样本的起始位置
    last = length - window_size
    if not starts or starts[-1] != last:
        starts.append(last)
    return starts   # list


class MaskedWellLogWindowDataset(Dataset):
    def __init__(
        self, stores, window_size=128, stride=32, channel_mask_prob=0.25,
        depth_mask_prob=0.75, depth_span=16, depth_spans=2,
    ):
        self.stores = stores
        self.window_size = int(window_size)
        self.channel_mask_prob = float(channel_mask_prob)   
        self.depth_mask_prob = float(depth_mask_prob)
        self.depth_span = int(depth_span)   # mask的长度
        self.depth_spans = int(depth_spans)  # mask的数量
        self.refs = [
            WindowRef(well, start)
            for well, store in stores.items()
            for start in window_starts(len(store["feat"]), self.window_size, stride)
        ]
        if not self.refs:
            raise ValueError("No pretraining windows were created; check well lengths and window size")

    # 掩码逻辑：
    # 对于每个样本，先生成一个随机数，若随机数小于channel_mask_prob，则进行channel mask，则随机选择一个channel全部mask掉；
    # 再生成一个随机数，如果随机数小于depth_mask_prob，则随机选择depth_spans个位置，每个位置的长度为depth_span，进行mask操作；
    # 最后如果没有任何mask，则随机选择一个channel和一个位置进行mask，确保至少有一个mask存在
    def _visible_mask(self, channels: int) -> torch.Tensor:
        visible = torch.ones(channels, self.window_size, dtype=torch.bool)   # (chanels, window_size)  # 初始化为全可见
        if random.random() < self.channel_mask_prob:  # channel mask 纵向mask
            visible[random.randrange(channels)] = False
        if random.random() < self.depth_mask_prob:   # depth mask 横向mask
            for _ in range(self.depth_spans):
                span = min(self.depth_span, self.window_size)  # mask长度
                start = random.randrange(self.window_size - span + 1)  # mask起始位置
                if random.random() < 0.7:
                    visible[:, start:start + span] = False   # 所有channel的同一位置都mask，mask长度为depth_span
                else:
                    visible[random.randrange(channels), start:start + span] = False   # 更弱mask，只mask一个channel的一段位置，mask长度为depth_span
        if visible.all():  # 兜底，确保样本有mask
            channel = random.randrange(channels)  # 随机选一个channel
            start = random.randrange(self.window_size - min(self.depth_span, self.window_size) + 1)   # 随机选一个起始深度
            visible[channel, start:start + min(self.depth_span, self.window_size)] = False
        # Never create an entirely invisible sample.
        if not visible.any():
            visible[:, :max(1, self.window_size // 8)] = True
        return visible

    def __len__(self):
        return len(self.refs)

    def __getitem__(self, index):
        ref = self.refs[index]
        feat = self.stores[ref.well]["feat"][ref.start:ref.start + self.window_size]   # 一个样本
        target = torch.from_numpy(feat.T.copy()).float()
        return target, self._visible_mask(target.shape[0])
