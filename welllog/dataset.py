from __future__ import annotations

from collections import Counter
import random
from typing import Dict, List, Optional, Tuple

import numpy as np
import torch
from openpyxl import load_workbook
from torch.utils.data import Dataset

# from welllog.emd_utils import attach_emd_spectra_to_store, build_emd_window_tensor
from welllog.input_utils import prepare_welllog_input
# from welllog.window_filter import is_pure_window


def _build_label_map_from_wells(wb, wells, label_col):
    """Collect sorted raw label strings from the given wells."""
    all_raw_labels = set()
    for well in wells:
        ws = wb[well]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        col_to_idx = {k: i for i, k in enumerate(header)}
        if label_col not in col_to_idx:
            continue
        label_idx = col_to_idx[label_col]
        for row in ws.iter_rows(min_row=2, values_only=True):
            v = row[label_idx]
            if v is not None and str(v).strip() != "":
                all_raw_labels.add(str(v).strip())
    return {lab: i for i, lab in enumerate(sorted(all_raw_labels))}


def _load_single_well(ws, feature_cols, label_col, depth_col, label_map):
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    col_to_idx = {k: i for i, k in enumerate(header)}
    for col in feature_cols + [label_col]:
        if col not in col_to_idx:
            raise ValueError(f"Column '{col}' not found in sheet '{ws.title}'.")

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    feat = np.zeros((len(rows), len(feature_cols)), dtype=np.float32)
    depths = []
    labels = []
    valid_mask = np.ones((len(rows),), dtype=bool)
    depth_idx = col_to_idx.get(depth_col)

    for i, row in enumerate(rows):  # 遍历每一行
        depths.append(row[depth_idx] if depth_idx is not None else None)
        for j, col in enumerate(feature_cols):  # 遍历每一列
            v = row[col_to_idx[col]]
            if v is None or str(v).strip() == "":
                feat[i, j] = np.nan
            else:
                feat[i, j] = float(v)
        raw_label = row[col_to_idx[label_col]]
        if raw_label is None or str(raw_label).strip() == "":
            valid_mask[i] = False
            labels.append(None)
        else:
            lab = str(raw_label).strip()
            if lab not in label_map:
                valid_mask[i] = False
                labels.append(None)
            else:
                labels.append(label_map[lab])

    for j in range(feat.shape[1]):  # 遍历每一列
        col_data = feat[:, j]
        col_median = np.nanmedian(col_data)
        if np.isnan(col_median):
            col_median = 0.0
        col_data[np.isnan(col_data)] = col_median
        feat[:, j] = col_data
    # 归一化
    # 每条曲线单独归一化
    feat_mean = feat.mean(axis=0, keepdims=True) # 计算每一列的均值
    feat_std = feat.std(axis=0, keepdims=True)
    feat_std[feat_std < 1e-6] = 1.0
    feat = (feat - feat_mean) / feat_std

    return {
        "feat": feat,
        "labels": labels,
        "depths": depths,
        "valid_mask": valid_mask,
    }


def load_welllog_store(args, is_train, label_map=None):
    """Load normalized per-well arrays from xlsx for train or val split."""
    feature_cols = [c.strip() for c in args.feature_cols.split(",")]
    label_col = args.label_col
    depth_col = args.depth_col
    label_map = dict(label_map) if label_map is not None else {}

    train_wells = [w.strip() for w in args.train_wells.split(",") if w.strip()]
    val_wells = [w.strip() for w in args.val_wells.split(",") if w.strip()]
    target_wells = train_wells if is_train else val_wells
    if len(target_wells) == 0:
        raise ValueError("No wells provided. Please set --train_wells/--val_wells.")

    wb = load_workbook(args.xlsx_path, data_only=True, read_only=True)
    available_wells = set(wb.sheetnames)
    for well in target_wells:
        if well not in available_wells:
            raise ValueError(f"Well '{well}' not found in xlsx sheets.")

    if not label_map:
        label_map = _build_label_map_from_wells(wb, target_wells, label_col)

    well_store = {}
    for well in target_wells:
        well_store[well] = _load_single_well(
            wb[well], feature_cols, label_col, depth_col, label_map
        )

    return target_wells, label_map, well_store


class WellLogSlidingWindowDataset(Dataset):
    """Well-log facies dataset using data-level sliding windows."""

    def __init__(self, args, is_train, label_map=None):
        super().__init__()
        self.well_input_mode = getattr(args, "well_input_mode", "sliding_window")
        self.input_size = args.input_size
        self.use_1d_conv = getattr(args, "use_1d_conv", False)
        self.task_mode = getattr(args, "task_mode", "classification")
        self.ignore_index = getattr(args, "ignore_index", -100)
        self.window_size = args.window_size
        self.window_require_pure = bool(getattr(args, "window_require_pure", False))
        if args.window_stride > 0:
            self.window_stride = args.window_stride
        elif self.task_mode == "segmentation":
            self.window_stride = max(args.window_size // 4, 1)
        else:
            self.window_stride = max(args.window_size // 2, 1)
        self.purity_thresh = getattr(args, "window_purity_thresh", 0.0) if is_train else 0.0
        self.purity_weight_power = getattr(args, "purity_weight_power", 0.0) if is_train else 0.0

        self.target_wells, self.label_map, self._well_store = load_welllog_store(
            args, is_train, label_map=label_map
        )
        self.class_counts = Counter()
        self.samples = []
        half = self.window_size // 2

        for well in self.target_wells:
            store = self._well_store[well]
            feat = store["feat"]
            labels = store["labels"]
            valid_mask = store["valid_mask"]
            n = len(feat)

            if n < self.window_size:
                continue
            if self.task_mode == "segmentation":
                start_iter = range(0, n - self.window_size + 1, self.window_stride)
            else:
                start_iter = (
                    center - half
                    for center in range(
                        half,
                        n - (self.window_size - half) + 1,
                        self.window_stride,
                    )
                )

            for start in start_iter:
                x_win = feat[start : start + self.window_size]

                if self.task_mode == "segmentation":
                    y_seq = [
                        labels[start + k]
                        if valid_mask[start + k] and labels[start + k] is not None
                        else self.ignore_index
                        for k in range(self.window_size)
                    ]
                    valid_labels_in_win = [y for y in y_seq if y != self.ignore_index]
                    if not valid_labels_in_win:
                        continue
                    self.samples.append((x_win, y_seq, 1.0))
                    self.class_counts.update(valid_labels_in_win)
                    continue

                center = start + half
                if not valid_mask[center]:
                    continue
                y = labels[center]

                valid_labels_in_win = [
                    labels[start + k]
                    for k in range(self.window_size)
                    if valid_mask[start + k] and labels[start + k] is not None
                ]
                if valid_labels_in_win:
                    purity = sum(1 for label in valid_labels_in_win if label == y) / len(valid_labels_in_win)
                else:
                    purity = 1.0

                if purity < self.purity_thresh:
                    continue

                sample_w = float(purity ** self.purity_weight_power) if self.purity_weight_power > 0.0 else 1.0
                self.samples.append((x_win, y, sample_w))
                self.class_counts[y] += 1

        if len(self.samples) == 0:
            raise ValueError("No valid samples built. Check well split and window size.")

    @property
    def inv_label_map(self):
        return {v: k for k, v in self.label_map.items()}

    def iter_dense(self, well):
        if well not in self._well_store:
            raise KeyError(f"Well '{well}' not found in this dataset split.")
        store = self._well_store[well]
        feat = store["feat"]
        labels = store["labels"]
        depths = store["depths"]
        valid_mask = store["valid_mask"]
        n = len(feat)
        half = self.window_size // 2
        if self.task_mode == "segmentation":
            for start in range(0, n - self.window_size + 1):
                x_win = feat[start : start + self.window_size]
                y_seq = [
                    labels[start + k]
                    if valid_mask[start + k] and labels[start + k] is not None
                    else self.ignore_index
                    for k in range(self.window_size)
                ]
                x = prepare_welllog_input(
                    x_win, self.input_size, self.use_1d_conv, task_mode=self.task_mode
                )
                yield depths[start : start + self.window_size], x, y_seq
            return

        for center in range(half, n - (self.window_size - half) + 1):
            if not valid_mask[center]:
                continue
            x_win = feat[center - half : center - half + self.window_size]
            x = prepare_welllog_input(
                x_win, self.input_size, self.use_1d_conv, task_mode=self.task_mode
            )
            yield depths[center], x, labels[center]

    def __len__(self):
        return len(self.samples)

    def __getitem__(self, idx):
        x_win, y, w = self.samples[idx]
        x = prepare_welllog_input(
            x_win, self.input_size, self.use_1d_conv, task_mode=self.task_mode
        )
        return x, torch.tensor(y, dtype=torch.long), torch.tensor(w, dtype=torch.float32)


class WellLogWholeWellDataset(Dataset):
    """One sample = one complete well section (variable length)."""

    def __init__(self, args, is_train, label_map=None):
        super().__init__()
        self.well_input_mode = "whole_well"
        self.is_train = is_train
        self.use_1d_conv = getattr(args, "use_1d_conv", False)
        self.task_mode = getattr(args, "task_mode", "segmentation")
        self.ignore_index = getattr(args, "ignore_index", -100)   # 将padding部分设为ignore_index
        self.max_seq_length = int(getattr(args, "whole_well_max_length", 0) or 0)  # 参数默认值0代表用全井段
        self.window_size = getattr(args, "window_size", 128)
        self.input_size = getattr(args, "input_size", self.window_size)

        self.target_wells, self.label_map, self._well_store = load_welllog_store(
            args, is_train, label_map=label_map
        )
        self.class_counts = Counter()
        self.well_names = []

        for well in self.target_wells: 
            store = self._well_store[well]
            labels = store["labels"]
            valid_mask = store["valid_mask"]  
            valid_labels = [
                labels[i]
                for i in range(len(labels))
                if valid_mask[i] and labels[i] is not None
            ]
            if not valid_labels:
                continue
            self.well_names.append(well)
            self.class_counts.update(valid_labels)

        if not self.well_names:
            raise ValueError("No wells with valid labels found for whole-well mode.")

    @property
    def inv_label_map(self):
        return {v: k for k, v in self.label_map.items()}

    def _labels_to_tensor(self, labels, valid_mask):
        y = []
        for i, label in enumerate(labels):
            if valid_mask[i] and label is not None:
                y.append(label)
            else:
                y.append(self.ignore_index)
        return torch.tensor(y, dtype=torch.long)

    def _maybe_crop(self, feat, labels, valid_mask):
        n = len(feat)  # 样本长度
        if self.max_seq_length <= 0 or n <= self.max_seq_length:
            return feat, labels, valid_mask  # 如果样本长度小于等于最大长度，则直接返回

        if self.is_train:
            start = random.randint(0, n - self.max_seq_length)
        else:
            start = 0
        end = start + self.max_seq_length
        return (
            feat[start:end],
            labels[start:end],
            valid_mask[start:end],
        )

    def __len__(self):
        return len(self.well_names)

    def __getitem__(self, idx):
        well = self.well_names[idx]
        store = self._well_store[well]
        feat, labels, valid_mask = self._maybe_crop(
            store["feat"], store["labels"], store["valid_mask"]
        )
        x = torch.from_numpy(feat.T.copy()).float()
        y = self._labels_to_tensor(labels, valid_mask)
        return x, y, torch.tensor(1.0, dtype=torch.float32)
