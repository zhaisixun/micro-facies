#!/usr/bin/env python3
"""
Plot probability density distributions of well-log curves (GR, CNL, DEN, etc.)
for all wells in an xlsx workbook (one sheet per well).

Example:
    python scripts/plot_welllog_density.py \\
        --xlsx_path ./facies-gr-diff0614-用GR-CNL-DEN.xlsx \\
        --output_dir ./outputs/welllog_density
"""

from __future__ import annotations

import argparse
import os
from typing import Dict, List

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import rcParams
from openpyxl import load_workbook
from scipy.stats import gaussian_kde


def _setup_sci_style() -> None:
    """Configure matplotlib for publication-quality figures."""
    rcParams.update(
        {
            "font.family": "serif",
            "font.serif": [
                "Times New Roman",
                "Noto Serif CJK SC",
                "Noto Serif CJK TC",
                "DejaVu Serif",
                "serif",
            ],
            "axes.unicode_minus": False,
            "font.size": 10,
            "axes.labelsize": 11,
            "axes.titlesize": 12,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "legend.fontsize": 7,
            "figure.dpi": 100,
            "savefig.dpi": 300,
            "savefig.bbox": "tight",
            "axes.linewidth": 0.8,
            "xtick.direction": "in",
            "ytick.direction": "in",
            "xtick.major.width": 0.8,
            "ytick.major.width": 0.8,
            "axes.grid": True,
            "grid.alpha": 0.25,
            "grid.linestyle": "--",
            "grid.linewidth": 0.5,
        }
    )


def _parse_feature_cols(feature_cols: str) -> List[str]:
    return [c.strip() for c in feature_cols.split(",") if c.strip()]


def _load_curve_values(
    xlsx_path: str,
    feature_cols: List[str],
) -> Dict[str, Dict[str, np.ndarray]]:
    """Return {well_name: {curve_name: 1d float array}}."""
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    well_data: Dict[str, Dict[str, np.ndarray]] = {}

    for well in wb.sheetnames:
        ws = wb[well]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        col_to_idx = {str(k): i for i, k in enumerate(header) if k is not None}

        missing = [col for col in feature_cols if col not in col_to_idx]
        if missing:
            raise ValueError(
                f"Columns {missing} not found in sheet '{well}'. "
                f"Available: {list(col_to_idx.keys())}"
            )

        curve_buffers = {col: [] for col in feature_cols}
        for row in ws.iter_rows(min_row=2, values_only=True):
            for col in feature_cols:
                v = row[col_to_idx[col]]
                if v is None or str(v).strip() == "":
                    continue
                try:
                    curve_buffers[col].append(float(v))
                except (TypeError, ValueError):
                    continue

        well_data[well] = {
            col: np.asarray(vals, dtype=np.float64)
            for col, vals in curve_buffers.items()
            if len(vals) > 0
        }

    wb.close()
    return well_data


def _kde_curve(values: np.ndarray, x_grid: np.ndarray) -> np.ndarray:
    kde = gaussian_kde(values)
    return kde(x_grid)


def _build_x_grid(all_vals: np.ndarray, kde_points: int) -> np.ndarray:
    x_min, x_max = np.percentile(all_vals, [0.5, 99.5])
    if x_max <= x_min:
        x_min, x_max = all_vals.min(), all_vals.max()
    pad = 0.05 * (x_max - x_min) if x_max > x_min else 1.0
    return np.linspace(x_min - pad, x_max + pad, kde_points)


def plot_density_distributions(
    well_data: Dict[str, Dict[str, np.ndarray]],
    feature_cols: List[str],
    output_dir: str,
    xlsx_basename: str,
    kde_points: int = 512,
) -> None:
    os.makedirs(output_dir, exist_ok=True)
    wells = sorted(well_data.keys())
    cmap = plt.get_cmap("tab20")
    colors = [cmap(i % 20) for i in range(len(wells))]

    for col in feature_cols:
        col_values = []
        for well in wells:
            vals = well_data[well].get(col)
            if vals is not None and vals.size > 0:
                col_values.append(vals)

        if not col_values:
            print(f"[skip] No valid data for column '{col}'.")
            continue

        all_vals = np.concatenate(col_values)
        x_grid = _build_x_grid(all_vals, kde_points)

        fig, ax = plt.subplots(figsize=(8, 5))
        plotted_wells = 0

        for well, color in zip(wells, colors):
            values = well_data[well].get(col)
            if values is None or values.size < 2:
                continue
            density = _kde_curve(values, x_grid)
            ax.plot(x_grid, density, color=color, linewidth=1.2, alpha=0.85, label=well)
            plotted_wells += 1

        ax.set_xlabel(col)
        ax.set_ylabel("Probability density")
        ax.set_title(f"{col} probability density — all wells ({plotted_wells} wells)")
        ax.legend(
            loc="center left",
            bbox_to_anchor=(1.02, 0.5),
            frameon=True,
            ncol=1,
        )

        fig.tight_layout()
        out_png = os.path.join(output_dir, f"{col}_density_all_wells.png")
        out_pdf = os.path.join(output_dir, f"{col}_density_all_wells.pdf")
        fig.savefig(out_png)
        fig.savefig(out_pdf)
        plt.close(fig)
        print(f"Saved: {out_png}")
        print(f"Saved: {out_pdf}")

    n_cols = len(feature_cols)
    fig, axes = plt.subplots(1, n_cols, figsize=(4.5 * n_cols, 4.5), squeeze=False)
    for ax, col in zip(axes[0], feature_cols):
        col_values = []
        for well in wells:
            vals = well_data[well].get(col)
            if vals is not None and vals.size > 0:
                col_values.append(vals)
        if not col_values:
            ax.set_visible(False)
            continue

        all_vals = np.concatenate(col_values)
        x_grid = _build_x_grid(all_vals, kde_points)

        for well, color in zip(wells, colors):
            values = well_data[well].get(col)
            if values is None or values.size < 2:
                continue
            density = _kde_curve(values, x_grid)
            ax.plot(x_grid, density, color=color, linewidth=1.0, alpha=0.8)

        ax.set_xlabel(col)
        ax.set_ylabel("Probability density")
        ax.set_title(col)

    handles = [
        plt.Line2D([0], [0], color=color, linewidth=1.2, label=well)
        for well, color in zip(wells, colors)
    ]
    fig.legend(
        handles=handles,
        loc="center left",
        bbox_to_anchor=(1.0, 0.5),
        fontsize=7,
        frameon=True,
    )
    fig.suptitle(f"Well-log density distributions — {xlsx_basename}", y=1.02)
    fig.tight_layout()
    combo_png = os.path.join(output_dir, "all_curves_density_all_wells.png")
    combo_pdf = os.path.join(output_dir, "all_curves_density_all_wells.pdf")
    fig.savefig(combo_png, bbox_inches="tight")
    fig.savefig(combo_pdf, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved: {combo_png}")
    print(f"Saved: {combo_pdf}")


def get_args_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Plot probability density of well-log curves for all wells."
    )
    p.add_argument(
        "--xlsx_path",
        default="./facies-gr-diff0614-用GR-CNL-DEN.xlsx",
        type=str,
        help="xlsx path (one sheet per well)",
    )
    p.add_argument(
        "--feature_cols",
        default="GR,CNL,DEN",
        type=str,
        help="comma-separated curve column names",
    )
    p.add_argument(
        "--output_dir",
        default="./outputs/welllog_density",
        type=str,
        help="directory to save figures",
    )
    p.add_argument(
        "--kde_points",
        default=512,
        type=int,
        help="number of points on the KDE evaluation grid",
    )
    return p


def main() -> None:
    args = get_args_parser().parse_args()
    _setup_sci_style()

    feature_cols = _parse_feature_cols(args.feature_cols)
    if not feature_cols:
        raise ValueError("--feature_cols must contain at least one column name.")

    xlsx_path = os.path.abspath(args.xlsx_path)
    if not os.path.isfile(xlsx_path):
        raise FileNotFoundError(f"xlsx not found: {xlsx_path}")

    well_data = _load_curve_values(xlsx_path, feature_cols)
    if not well_data:
        raise ValueError(f"No wells loaded from '{xlsx_path}'.")

    print(f"Loaded {len(well_data)} wells from {xlsx_path}")
    plot_density_distributions(
        well_data=well_data,
        feature_cols=feature_cols,
        output_dir=args.output_dir,
        xlsx_basename=os.path.basename(xlsx_path),
        kde_points=args.kde_points,
    )


if __name__ == "__main__":
    main()
