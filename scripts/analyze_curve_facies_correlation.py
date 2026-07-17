#!/usr/bin/env python3
"""
Analyze well-log curve vs sedimentary microfacies (facies) relationships.

Reads xlsx (one sheet per well), computes separability metrics, plots
correlation heatmaps and facies-grouped boxplots, and suggests candidate
feature_cols for run_hyperparam_experiments.py ablation.

Example:
    python scripts/analyze_curve_facies_correlation.py \\
        --xlsx_path ./facies-gr-diff0614-用GR-CNL-DEN.xlsx \\
        --output_dir outputs/curve_facies_correlation

Auto-discover all numeric curve columns:
    python scripts/analyze_curve_facies_correlation.py \\
        --xlsx_path ./facies-gr-diff0614-用GR-CNL-DEN.xlsx \\
        --feature_cols auto \\
        --output_dir outputs/curve_facies_correlation
"""

from __future__ import annotations

import argparse
import csv
import os
from collections import Counter
from typing import Dict, List, Optional, Sequence, Tuple

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import rcParams
from openpyxl import load_workbook
from scipy.stats import kruskal, spearmanr
from sklearn.feature_selection import f_classif, mutual_info_classif


# Colorblind-friendly palette (Tol bright)
CLASS_COLORS = ["#4477AA", "#EE6677", "#228833", "#CCBB44", "#AA3377", "#66CCEE", "#BBBBBB"]

DEFAULT_EXCLUDE_COLS = {"亚相", "微相", "Well", "WELL", "well", "井名"}


def _setup_sci_style() -> None:
    rcParams.update(
        {
            "font.family": "serif",
            "font.serif": [
                "Times New Roman",
                "Noto Serif CJK SC",
                "Noto Serif CJK TC",
                "DejaVu Serif",
                "serif",
            ],
            "axes.unicode_minus": False,
            "font.size": 10,
            "axes.labelsize": 11,
            "axes.titlesize": 12,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "legend.fontsize": 9,
            "figure.dpi": 100,
            "savefig.dpi": 300,
            "savefig.bbox": "tight",
            "axes.linewidth": 0.8,
            "xtick.direction": "in",
            "ytick.direction": "in",
            "xtick.major.width": 0.8,
            "ytick.major.width": 0.8,
            "axes.grid": True,
            "grid.alpha": 0.25,
            "grid.linestyle": "--",
            "grid.linewidth": 0.5,
        }
    )


def _parse_cols(raw: str) -> List[str]:
    return [c.strip() for c in raw.split(",") if c.strip()]


def _build_label_map(wb, wells: Sequence[str], label_col: str) -> Dict[str, int]:
    all_raw_labels = set()
    for well in wells:
        ws = wb[well]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        col_to_idx = {str(k): i for i, k in enumerate(header) if k is not None}
        if label_col not in col_to_idx:
            continue
        label_idx = col_to_idx[label_col]
        for row in ws.iter_rows(min_row=2, values_only=True):
            v = row[label_idx]
            if v is not None and str(v).strip() != "":
                all_raw_labels.add(str(v).strip())
    if not all_raw_labels:
        raise ValueError(f"No labels found in column '{label_col}'.")
    return {lab: i for i, lab in enumerate(sorted(all_raw_labels))}


def _is_numeric_value(v) -> bool:
    if v is None or str(v).strip() == "":
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _discover_numeric_columns(
    wb,
    wells: Sequence[str],
    depth_col: str,
    label_col: str,
    exclude_cols: Sequence[str],
    min_valid_ratio: float,
) -> List[str]:
    exclude = set(exclude_cols) | DEFAULT_EXCLUDE_COLS | {depth_col, label_col}
    common_cols: Optional[set[str]] = None
    counts: Counter[str] = Counter()
    numeric_counts: Counter[str] = Counter()

    for well in wells:
        ws = wb[well]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        col_names = [str(k) for k in header if k is not None]
        well_cols = set(col_names)
        common_cols = well_cols if common_cols is None else common_cols & well_cols

        for col in col_names:
            if col in exclude:
                continue
            col_idx = col_names.index(col)
            for row in ws.iter_rows(min_row=2, values_only=True):
                v = row[col_idx]
                if v is None or str(v).strip() == "":
                    continue
                counts[col] += 1
                if _is_numeric_value(v):
                    numeric_counts[col] += 1

    if common_cols is None:
        raise ValueError("No wells available for column discovery.")

    discovered: List[str] = []
    for col in sorted(common_cols):
        if col in exclude:
            continue
        total = counts[col]
        if total == 0:
            continue
        ratio = numeric_counts[col] / total
        if ratio >= min_valid_ratio:
            discovered.append(col)
    if not discovered:
        raise ValueError(
            "No numeric curve columns discovered across all wells. "
            "Pass --feature_cols explicitly or reduce --min_valid_ratio."
        )
    return discovered


def _resolve_feature_cols(
    wb,
    wells: Sequence[str],
    feature_cols_arg: str,
    depth_col: str,
    label_col: str,
    exclude_cols: Sequence[str],
    min_valid_ratio: float,
) -> List[str]:
    if feature_cols_arg.strip().lower() == "auto":
        return _discover_numeric_columns(
            wb, wells, depth_col, label_col, exclude_cols, min_valid_ratio
        )
    cols = _parse_cols(feature_cols_arg)
    if not cols:
        raise ValueError("--feature_cols must be 'auto' or a comma-separated column list.")
    return cols


def _load_samples(
    xlsx_path: str,
    feature_cols: List[str],
    label_col: str,
    depth_col: str,
    wells: Optional[List[str]] = None,
    zscore: bool = True,
) -> Tuple[np.ndarray, np.ndarray, Dict[str, int], Dict[int, str], List[str]]:
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    if wells is None:
        wells = list(wb.sheetnames)
    if not wells:
        raise ValueError(f"No sheets found in '{xlsx_path}'.")

    label_map = _build_label_map(wb, wells, label_col)
    inv_label_map = {v: k for k, v in label_map.items()}

    xs: List[np.ndarray] = []
    ys: List[np.ndarray] = []

    for well in wells:
        if well not in wb.sheetnames:
            raise ValueError(f"Well '{well}' not found in xlsx sheets.")
        ws = wb[well]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        col_to_idx = {str(k): i for i, k in enumerate(header) if k is not None}

        for col in feature_cols + [label_col]:
            if col not in col_to_idx:
                raise ValueError(
                    f"Column '{col}' not found in sheet '{well}'. "
                    f"Available: {list(col_to_idx.keys())}"
                )

        feat_idx = [col_to_idx[c] for c in feature_cols]
        label_idx = col_to_idx[label_col]

        well_x: List[List[float]] = []
        well_y: List[int] = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_label = row[label_idx]
            if raw_label is None or str(raw_label).strip() == "":
                continue
            lab = str(raw_label).strip()
            if lab not in label_map:
                continue

            values: List[float] = []
            skip = False
            for fi in feat_idx:
                v = row[fi]
                if v is None or str(v).strip() == "":
                    skip = True
                    break
                values.append(float(v))
            if skip:
                continue

            well_x.append(values)
            well_y.append(label_map[lab])

        if well_x:
            xs.append(np.asarray(well_x, dtype=np.float64))
            ys.append(np.asarray(well_y, dtype=np.int64))

    wb.close()

    if not xs:
        raise ValueError("No valid labeled samples loaded. Check columns and facies labels.")

    X = np.vstack(xs)
    y = np.concatenate(ys)

    if zscore:
        for j in range(X.shape[1]):
            col = X[:, j]
            col_median = np.nanmedian(col)
            if np.isnan(col_median):
                col_median = 0.0
            col[np.isnan(col)] = col_median
            mean = col.mean()
            std = col.std()
            if std < 1e-6:
                std = 1.0
            X[:, j] = (col - mean) / std

    return X, y, label_map, inv_label_map, wells


def _compute_curve_scores(
    X: np.ndarray,
    y: np.ndarray,
    feature_cols: List[str],
    seed: int,
) -> Dict[str, Dict[str, float]]:
    f_vals, _ = f_classif(X, y)
    mi_vals = mutual_info_classif(X, y, discrete_features=False, random_state=seed)

    scores: Dict[str, Dict[str, float]] = {}
    for j, col in enumerate(feature_cols):
        col_x = X[:, j]
        pearson_r = float(np.corrcoef(col_x, y)[0, 1])
        if np.isnan(pearson_r):
            pearson_r = 0.0
        sp_r, _ = spearmanr(col_x, y)
        sp_r = float(sp_r) if not np.isnan(sp_r) else 0.0

        groups = [col_x[y == cls] for cls in sorted(set(y.tolist()))]
        groups = [g for g in groups if g.size > 0]
        if len(groups) >= 2:
            kw_h, kw_p = kruskal(*groups)
        else:
            kw_h, kw_p = 0.0, 1.0

        scores[col] = {
            "pearson_r": pearson_r,
            "pearson_abs": abs(pearson_r),
            "spearman_r": sp_r,
            "spearman_abs": abs(sp_r),
            "anova_f": float(f_vals[j]),
            "mutual_info": float(mi_vals[j]),
            "kruskal_h": float(kw_h),
            "kruskal_p": float(kw_p),
        }

    metrics_for_rank = ["pearson_abs", "spearman_abs", "anova_f", "mutual_info", "kruskal_h"]
    ranks = {m: {} for m in metrics_for_rank}
    for metric in metrics_for_rank:
        ordered = sorted(feature_cols, key=lambda c: scores[c][metric], reverse=True)
        for rank, col in enumerate(ordered, start=1):
            ranks[metric][col] = rank

    for col in feature_cols:
        avg_rank = sum(ranks[m][col] for m in metrics_for_rank) / len(metrics_for_rank)
        scores[col]["composite_rank"] = avg_rank
        scores[col]["composite_score"] = 1.0 / avg_rank

    return scores


def _curve_corr_matrix(X: np.ndarray) -> np.ndarray:
    corr = np.corrcoef(X, rowvar=False)
    corr = np.nan_to_num(corr, nan=0.0, posinf=0.0, neginf=0.0)
    np.fill_diagonal(corr, 1.0)
    return corr


def _select_top_k(
    feature_cols: List[str],
    scores: Dict[str, Dict[str, float]],
    metric: str,
    k: int,
) -> List[str]:
    ranked = sorted(feature_cols, key=lambda c: scores[c][metric], reverse=True)
    return ranked[:k]


def _select_diverse_top_k(
    feature_cols: List[str],
    scores: Dict[str, Dict[str, float]],
    corr: np.ndarray,
    k: int,
    max_corr: float,
) -> List[str]:
    ranked = sorted(feature_cols, key=lambda c: scores[c]["composite_score"], reverse=True)
    col_to_idx = {c: i for i, c in enumerate(feature_cols)}
    selected: List[str] = []

    for col in ranked:
        if len(selected) >= k:
            break
        if all(abs(corr[col_to_idx[col], col_to_idx[s]]) <= max_corr for s in selected):
            selected.append(col)

    for col in ranked:
        if len(selected) >= k:
            break
        if col not in selected:
            selected.append(col)
    return selected[:k]


def _unique_feature_sets(candidates: List[List[str]]) -> List[str]:
    seen = set()
    out: List[str] = []
    for cols in candidates:
        key = tuple(cols)
        if key in seen:
            continue
        seen.add(key)
        out.append(",".join(cols))
    return out


def _recommend_feature_sets(
    feature_cols: List[str],
    scores: Dict[str, Dict[str, float]],
    corr: np.ndarray,
    top_k: int,
    max_corr: float,
    baseline_cols: Optional[List[str]] = None,
) -> Dict[str, List[str]]:
    rec: Dict[str, List[str]] = {
        "top_k_pearson": _select_top_k(feature_cols, scores, "pearson_abs", top_k),
        "top_k_spearman": _select_top_k(feature_cols, scores, "spearman_abs", top_k),
        "top_k_anova_f": _select_top_k(feature_cols, scores, "anova_f", top_k),
        "top_k_mutual_info": _select_top_k(feature_cols, scores, "mutual_info", top_k),
        "top_k_composite": _select_top_k(feature_cols, scores, "composite_score", top_k),
        "top_k_diverse": _select_diverse_top_k(feature_cols, scores, corr, top_k, max_corr),
    }
    if baseline_cols:
        baseline = [c for c in baseline_cols if c in feature_cols]
        if baseline:
            rec["baseline"] = baseline
    return rec


def _print_class_counts(y: np.ndarray, inv_label_map: Dict[int, str]) -> None:
    counts = Counter(int(v) for v in y)
    print("\n=== Facies class counts ===")
    for cls in sorted(counts):
        name = inv_label_map.get(cls, str(cls))
        print(f"  class {cls} ({name}): {counts[cls]}")
    print(f"  total: {len(y)}")


def _save_scores_csv(
    path: str,
    feature_cols: List[str],
    scores: Dict[str, Dict[str, float]],
) -> None:
    fieldnames = [
        "curve",
        "pearson_r",
        "pearson_abs",
        "spearman_r",
        "spearman_abs",
        "anova_f",
        "mutual_info",
        "kruskal_h",
        "kruskal_p",
        "composite_rank",
        "composite_score",
    ]
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        ordered = sorted(feature_cols, key=lambda c: scores[c]["composite_rank"])
        for col in ordered:
            row = {"curve": col, **scores[col]}
            writer.writerow(row)
    print(f"Saved scores -> {path}")


def _plot_correlation_heatmap(
    corr: np.ndarray,
    feature_cols: List[str],
    scores: Dict[str, Dict[str, float]],
    output_path: str,
) -> None:
    n = len(feature_cols)
    fig_h = max(4.5, 0.45 * n + 2.0)
    fig, axes = plt.subplots(1, 2, figsize=(fig_h + 4.5, fig_h))

    im = axes[0].imshow(corr, vmin=-1.0, vmax=1.0, cmap="RdBu_r", aspect="auto")
    axes[0].set_xticks(range(n))
    axes[0].set_yticks(range(n))
    axes[0].set_xticklabels(feature_cols, rotation=45, ha="right")
    axes[0].set_yticklabels(feature_cols)
    axes[0].set_title("Curve–curve Pearson correlation")
    cbar = fig.colorbar(im, ax=axes[0], fraction=0.046, pad=0.04)
    cbar.set_label("Pearson r")

    for i in range(n):
        for j in range(n):
            axes[0].text(j, i, f"{corr[i, j]:.2f}", ha="center", va="center", fontsize=7)

    metric_names = ["|Pearson|", "|Spearman|", "ANOVA F", "Mutual info", "Kruskal H"]
    metric_keys = ["pearson_abs", "spearman_abs", "anova_f", "mutual_info", "kruskal_h"]
    x = np.arange(len(feature_cols))
    width = 0.15
    for idx, (mname, mkey) in enumerate(zip(metric_names, metric_keys)):
        vals = np.array([scores[c][mkey] for c in feature_cols], dtype=np.float64)
        if mkey in ("pearson_abs", "spearman_abs"):
            norm = vals
        else:
            vmax = vals.max() if vals.max() > 0 else 1.0
            norm = vals / vmax
        axes[1].bar(x + (idx - 2) * width, norm, width=width, label=mname, alpha=0.9)

    axes[1].set_xticks(x)
    axes[1].set_xticklabels(feature_cols, rotation=45, ha="right")
    axes[1].set_ylabel("Normalized score")
    axes[1].set_title("Curve–facies separability (normalized)")
    axes[1].legend(loc="best", fontsize=8, frameon=True)

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved figure -> {output_path}")


def _plot_boxplots_by_facies(
    X: np.ndarray,
    y: np.ndarray,
    feature_cols: List[str],
    inv_label_map: Dict[int, str],
    scores: Dict[str, Dict[str, float]],
    output_path: str,
    max_curves: int,
) -> None:
    ordered_cols = sorted(feature_cols, key=lambda c: scores[c]["composite_rank"])[:max_curves]
    n_cols = len(ordered_cols)
    n_rows = int(np.ceil(n_cols / 3))
    fig, axes = plt.subplots(n_rows, 3, figsize=(12, 3.8 * n_rows), squeeze=False)

    classes = sorted(set(int(v) for v in y))
    for ax_idx, col in enumerate(ordered_cols):
        r, c = divmod(ax_idx, 3)
        ax = axes[r][c]
        j = feature_cols.index(col)
        data = [X[y == cls, j] for cls in classes]
        positions = np.arange(len(classes))
        bp = ax.boxplot(
            data,
            positions=positions,
            widths=0.55,
            patch_artist=True,
            showfliers=False,
            medianprops={"color": "black", "linewidth": 1.2},
        )
        for patch, cls in zip(bp["boxes"], classes):
            patch.set_facecolor(CLASS_COLORS[cls % len(CLASS_COLORS)])
            patch.set_alpha(0.75)
        ax.set_xticks(positions)
        ax.set_xticklabels([inv_label_map.get(cls, str(cls)) for cls in classes], rotation=20, ha="right")
        ax.set_title(f"{col}  (rank={scores[col]['composite_rank']:.1f})")
        ax.set_ylabel("z-score" if X.std(axis=0).mean() > 0 else "value")

    for ax_idx in range(n_cols, n_rows * 3):
        r, c = divmod(ax_idx, 3)
        axes[r][c].set_visible(False)

    fig.suptitle("Curve distributions by facies (top curves by composite rank)", y=1.01)
    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved figure -> {output_path}")


def _plot_separability_bars(
    feature_cols: List[str],
    scores: Dict[str, Dict[str, float]],
    output_path: str,
) -> None:
    ordered = sorted(feature_cols, key=lambda c: scores[c]["composite_rank"])
    composite = [scores[c]["composite_score"] for c in ordered]

    fig, ax = plt.subplots(figsize=(max(6, 0.55 * len(ordered)), 4.5))
    y_pos = np.arange(len(ordered))
    ax.barh(y_pos, composite, color="#4477AA", alpha=0.85)
    ax.set_yticks(y_pos)
    ax.set_yticklabels(ordered)
    ax.invert_yaxis()
    ax.set_xlabel("Composite score (1 / mean rank; higher = better)")
    ax.set_title("Overall curve–facies separability ranking")
    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved figure -> {output_path}")


def _write_recommendations(
    output_dir: str,
    recommendations: Dict[str, List[str]],
    feature_cols: List[str],
    feature_cols_list: List[str],
    scores: Dict[str, Dict[str, float]],
    top_k: int,
    max_corr: float,
) -> None:
    txt_path = os.path.join(output_dir, "recommendations.txt")
    md_path = os.path.join(output_dir, "recommendations.md")

    lines = [
        "Curve–facies correlation analysis recommendations",
        "=" * 60,
        f"top_k={top_k}, diverse max_corr={max_corr}",
        "",
        "Per-curve composite ranking (best first):",
    ]
    for col in sorted(feature_cols, key=lambda c: scores[c]["composite_rank"]):
        s = scores[col]
        lines.append(
            f"  {col:12s}  rank={s['composite_rank']:.2f}  "
            f"|r|={s['pearson_abs']:.4f}  MI={s['mutual_info']:.4f}  F={s['anova_f']:.2f}"
        )

    lines.append("\nSuggested feature_cols sets:")
    for name, cols in recommendations.items():
        lines.append(f"  [{name}]  {','.join(cols)}")

    lines.append("\nFor run_hyperparam_experiments.py:")
    lines.append(f'  --feature_cols_list "{("|".join(feature_cols_list))}"')
    lines.append("\nNext step: run OFAT ablation and compare mIoU / macro F1 on val set.")

    with open(txt_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    print(f"Saved recommendations -> {txt_path}")

    md_lines = [
        "# Curve–facies correlation recommendations",
        "",
        f"- top_k: **{top_k}**",
        f"- diverse selection max |corr|: **{max_corr}**",
        "",
        "## Composite ranking",
        "",
        "| Curve | Rank | |Pearson| | MI | ANOVA F |",
        "|-------|------|----------|-----|---------|",
    ]
    for col in sorted(feature_cols, key=lambda c: scores[c]["composite_rank"]):
        s = scores[col]
        md_lines.append(
            f"| {col} | {s['composite_rank']:.2f} | {s['pearson_abs']:.4f} | "
            f"{s['mutual_info']:.4f} | {s['anova_f']:.2f} |"
        )

    md_lines.extend(["", "## Suggested feature sets", ""])
    for name, cols in recommendations.items():
        md_lines.append(f"- **{name}**: `{','.join(cols)}`")

    md_lines.extend(
        [
            "",
            "## Ablation command snippet",
            "",
            "```bash",
            "python scripts/run_hyperparam_experiments.py \\",
            f'  --feature_cols_list "{("|".join(feature_cols_list))}" \\',
            "  ... # other args",
            "```",
            "",
            "> Statistical ranking is exploratory; final selection should use validation mIoU / macro F1.",
        ]
    )

    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(md_lines) + "\n")
    print(f"Saved recommendations -> {md_path}")


def run_analysis(args: argparse.Namespace) -> None:
    _setup_sci_style()
    os.makedirs(args.output_dir, exist_ok=True)

    xlsx_path = os.path.abspath(args.xlsx_path)
    if not os.path.isfile(xlsx_path):
        raise FileNotFoundError(f"xlsx not found: {xlsx_path}")

    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    wells = _parse_cols(args.wells) if args.wells.strip() else list(wb.sheetnames)
    exclude_cols = _parse_cols(args.exclude_cols)
    feature_cols = _resolve_feature_cols(
        wb,
        wells,
        args.feature_cols,
        args.depth_col,
        args.label_col,
        exclude_cols,
        args.min_valid_ratio,
    )
    wb.close()

    print(f"xlsx: {xlsx_path}")
    print(f"wells ({len(wells)}): {', '.join(wells)}")
    print(f"curves ({len(feature_cols)}): {', '.join(feature_cols)}")

    X, y, label_map, inv_label_map, _ = _load_samples(
        xlsx_path=xlsx_path,
        feature_cols=feature_cols,
        label_col=args.label_col,
        depth_col=args.depth_col,
        wells=wells,
        zscore=not args.no_zscore,
    )
    print(f"loaded samples: X={X.shape}, y={y.shape}")
    print(f"label_map: {label_map}")
    _print_class_counts(y, inv_label_map)

    scores = _compute_curve_scores(X, y, feature_cols, seed=args.seed)
    corr = _curve_corr_matrix(X)

    baseline_cols = _parse_cols(args.baseline_cols) if args.baseline_cols else _parse_cols("GR,CNL,DEN")
    recommendations = _recommend_feature_sets(
        feature_cols, scores, corr, args.top_k, args.max_corr, baseline_cols
    )

    candidate_lists = list(recommendations.values())
    feature_cols_list = _unique_feature_sets(candidate_lists)

    _save_scores_csv(os.path.join(args.output_dir, "curve_facies_scores.csv"), feature_cols, scores)
    _plot_correlation_heatmap(
        corr,
        feature_cols,
        scores,
        os.path.join(args.output_dir, "correlation_heatmap.png"),
    )
    _plot_boxplots_by_facies(
        X,
        y,
        feature_cols,
        inv_label_map,
        scores,
        os.path.join(args.output_dir, "curve_facies_boxplots.png"),
        max_curves=args.max_boxplot_curves,
    )
    _plot_separability_bars(
        feature_cols,
        scores,
        os.path.join(args.output_dir, "curve_separability_ranking.png"),
    )
    _write_recommendations(
        args.output_dir,
        recommendations,
        feature_cols,
        feature_cols_list,
        scores,
        args.top_k,
        args.max_corr,
    )

    print("\nAnalysis complete.")


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Analyze well-log curve vs facies correlation and suggest feature_cols."
    )
    p.add_argument(
        "--xlsx_path",
        default="./facies-gr-diff0614-用GR-CNL-DEN.xlsx",
        help="xlsx path (one sheet per well)",
    )
    p.add_argument(
        "--feature_cols",
        default="auto",
        help="Comma-separated curve columns, or 'auto' to discover numeric columns",
    )
    p.add_argument("--label_col", default="facies", help="Facies label column")
    p.add_argument("--depth_col", default="DEPT", help="Depth column (excluded from analysis)")
    p.add_argument(
        "--wells",
        default="",
        help="Comma-separated well sheet names; empty = all sheets",
    )
    p.add_argument(
        "--exclude_cols",
        default="亚相,微相",
        help="Extra non-curve columns to exclude when --feature_cols auto",
    )
    p.add_argument(
        "--min_valid_ratio",
        type=float,
        default=0.95,
        help="Min fraction of numeric values for auto column discovery",
    )
    p.add_argument("--output_dir", default="outputs/curve_facies_correlation", help="Output directory")
    p.add_argument("--top_k", type=int, default=3, help="Number of curves in each recommendation set")
    p.add_argument(
        "--max_corr",
        type=float,
        default=0.85,
        help="Max |corr| between curves in diverse top-k selection",
    )
    p.add_argument(
        "--baseline_cols",
        default="GR,CNL,DEN",
        help="Current baseline feature set to include in recommendations",
    )
    p.add_argument(
        "--max_boxplot_curves",
        type=int,
        default=9,
        help="Max number of curves in facies boxplot grid",
    )
    p.add_argument("--no_zscore", action="store_true", help="Skip per-curve z-score normalization")
    p.add_argument("--seed", type=int, default=42, help="Random seed for mutual information")
    return p.parse_args()


def main() -> None:
    run_analysis(parse_args())


if __name__ == "__main__":
    main()
