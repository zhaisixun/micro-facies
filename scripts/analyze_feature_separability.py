#!/usr/bin/env python3
"""
Analyze separability of sedimentary microfacies in original well-log feature space.

Usage (from npy files):
    python scripts/analyze_feature_separability.py \\
        --x_path data/X.npy --y_path data/y.npy --output_dir outputs/separability

Optional: export npy from xlsx first:
    python scripts/analyze_feature_separability.py \\
        --export_xlsx ./facies-gr-diff0607-用GR-CNL-DEN.xlsx \\
        --feature_cols GR,CNL,DEN --label_col facies \\
        --x_path data/X.npy --y_path data/y.npy
"""

from __future__ import annotations

import argparse
import os
from collections import Counter
from typing import Dict, List, Optional, Tuple

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import rcParams
from openpyxl import load_workbook
from sklearn.decomposition import PCA
from sklearn.manifold import TSNE
from sklearn.preprocessing import StandardScaler


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

CLASS_LABELS: Dict[int, str] = {
    0: "0",
    1: "1-河口坝",
    2: "2-席状砂",
    3: "3",
}

# Colorblind-friendly palette (Tol bright)
CLASS_COLORS: Dict[int, str] = {
    0: "#4477AA",
    1: "#EE6677",
    2: "#228833",
    3: "#CCBB44",
}


def _setup_sci_style() -> None:
    """Configure matplotlib for publication-quality figures."""
    rcParams.update(
        {
            "font.family": "serif",
            "font.serif": [
                "Times New Roman",
                "Noto Serif CJK SC",
                "Noto Serif CJK TC",
                "DejaVu Serif",
                "serif",
            ],
            "axes.unicode_minus": False,
            "font.size": 10,
            "axes.labelsize": 11,
            "axes.titlesize": 12,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "legend.fontsize": 9,
            "figure.dpi": 100,
            "savefig.dpi": 300,
            "savefig.bbox": "tight",
            "axes.linewidth": 0.8,
            "xtick.direction": "in",
            "ytick.direction": "in",
            "xtick.major.width": 0.8,
            "ytick.major.width": 0.8,
            "axes.grid": True,
            "grid.alpha": 0.25,
            "grid.linestyle": "--",
            "grid.linewidth": 0.5,
        }
    )


def export_xlsx_to_npy(
    xlsx_path: str,
    x_path: str,
    y_path: str,
    feature_cols: List[str],
    label_col: str = "facies",
) -> Tuple[np.ndarray, np.ndarray]:
    """Export all wells from xlsx into X (N, C) and y (N,) npy arrays."""
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    xs: List[np.ndarray] = []
    ys: List[np.ndarray] = []

    for well in wb.sheetnames:
        ws = wb[well]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        col_to_idx = {name: i for i, name in enumerate(header)}

        for col in feature_cols + [label_col]:
            if col not in col_to_idx:
                raise ValueError(f"Column '{col}' not found in sheet '{well}'.")

        feat_idx = [col_to_idx[c] for c in feature_cols]
        label_idx = col_to_idx[label_col]

        well_x: List[List[float]] = []
        well_y: List[int] = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_label = row[label_idx]
            if raw_label is None or str(raw_label).strip() == "":
                continue

            try:
                label = int(str(raw_label).strip())
            except ValueError as exc:
                raise ValueError(
                    f"Invalid label '{raw_label}' in well '{well}'. Expected integer class id."
                ) from exc

            values: List[float] = []
            skip = False
            for fi in feat_idx:
                v = row[fi]
                if v is None or str(v).strip() == "":
                    skip = True
                    break
                values.append(float(v))
            if skip:
                continue

            well_x.append(values)
            well_y.append(label)

        if well_x:
            xs.append(np.asarray(well_x, dtype=np.float32))
            ys.append(np.asarray(well_y, dtype=np.int64))

    if not xs:
        raise ValueError("No valid samples exported. Check xlsx content and column names.")

    X = np.vstack(xs)
    y = np.concatenate(ys)

    os.makedirs(os.path.dirname(os.path.abspath(x_path)) or ".", exist_ok=True)
    os.makedirs(os.path.dirname(os.path.abspath(y_path)) or ".", exist_ok=True)
    np.save(x_path, X)
    np.save(y_path, y)
    print(f"Exported X {X.shape} -> {x_path}")
    print(f"Exported y {y.shape} -> {y_path}")
    return X, y


def load_data(x_path: str, y_path: str) -> Tuple[np.ndarray, np.ndarray]:
    X = np.load(x_path)
    y = np.load(y_path)

    if X.ndim != 2:
        raise ValueError(f"X must be 2D (N, C), got shape {X.shape}")
    if y.ndim != 1:
        y = y.reshape(-1)
    if X.shape[0] != y.shape[0]:
        raise ValueError(f"Sample size mismatch: X={X.shape[0]}, y={y.shape[0]}")
    return X.astype(np.float64), y.astype(np.int64)


def print_class_counts(y: np.ndarray) -> None:
    counts = Counter(int(v) for v in y)
    print("\n=== Class counts (all samples) ===")
    for cls in sorted(counts):
        name = CLASS_LABELS.get(cls, str(cls))
        print(f"  class {cls} ({name}): {counts[cls]}")
    print(f"  total: {len(y)}")


def subsample(
    X: np.ndarray, y: np.ndarray, max_samples: int, seed: int
) -> Tuple[np.ndarray, np.ndarray]:
    n = X.shape[0]
    if n <= max_samples:
        return X, y
    rng = np.random.default_rng(seed)
    idx = rng.choice(n, size=max_samples, replace=False)
    return X[idx], y[idx]


def plot_embedding(
    coords: np.ndarray,
    y: np.ndarray,
    title: str,
    xlabel: str,
    ylabel: str,
    save_path: str,
) -> None:
    fig, ax = plt.subplots(figsize=(5.5, 4.5))

    classes = sorted(set(int(v) for v in y))
    for cls in classes:
        mask = y == cls
        color = CLASS_COLORS.get(cls, None)
        label = CLASS_LABELS.get(cls, str(cls))
        ax.scatter(
            coords[mask, 0],
            coords[mask, 1],
            c=color,
            label=label,
            s=14,
            alpha=0.75,
            edgecolors="none",
        )

    ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    ax.legend(
        title="Facies",
        loc="best",
        frameon=True,
        framealpha=0.9,
        edgecolor="0.8",
    )
    fig.tight_layout()
    os.makedirs(os.path.dirname(os.path.abspath(save_path)) or ".", exist_ok=True)
    fig.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    print(f"Saved figure -> {save_path}")


def run_analysis(
    X: np.ndarray,
    y: np.ndarray,
    output_dir: str,
    max_samples: int = 10000,
    seed: int = 42,
    tsne_perplexity: float = 30.0,
) -> None:
    _setup_sci_style()
    os.makedirs(output_dir, exist_ok=True)

    print_class_counts(y)

    # Standardize on all samples, then subsample for visualization
    scaler = StandardScaler()
    X_scaled = scaler.fit_transform(X)
    X_vis, y_vis = subsample(X_scaled, y, max_samples=max_samples, seed=seed)
    print(f"\nVisualization subset: {X_vis.shape[0]} / {X.shape[0]} samples")

    # PCA (fit on visualization subset for consistency with plotted points)
    pca = PCA(n_components=2, random_state=seed)
    X_pca = pca.fit_transform(X_vis)
    evr = pca.explained_variance_ratio_
    print("\n=== PCA explained variance ratio ===")
    print(f"  PC1: {evr[0]:.4f} ({evr[0] * 100:.2f}%)")
    print(f"  PC2: {evr[1]:.4f} ({evr[1] * 100:.2f}%)")
    print(f"  PC1+PC2 cumulative: {evr.sum():.4f} ({evr.sum() * 100:.2f}%)")

    plot_embedding(
        X_pca,
        y_vis,
        title="PCA of Original Well-Log Features",
        xlabel=f"PC1 ({evr[0] * 100:.1f}%)",
        ylabel=f"PC2 ({evr[1] * 100:.1f}%)",
        save_path=os.path.join(output_dir, "pca_original_features.png"),
    )

    # t-SNE
    n_vis = X_vis.shape[0]
    perplexity = min(tsne_perplexity, max(5.0, (n_vis - 1) / 3.0))
    print(f"\nRunning t-SNE (n={n_vis}, perplexity={perplexity:.1f}) ...")
    tsne = TSNE(
        n_components=2,
        perplexity=perplexity,
        random_state=seed,
        init="pca",
        learning_rate="auto",
        n_iter=1000,
    )
    X_tsne = tsne.fit_transform(X_vis)

    plot_embedding(
        X_tsne,
        y_vis,
        title="t-SNE of Original Well-Log Features",
        xlabel="t-SNE dimension 1",
        ylabel="t-SNE dimension 2",
        save_path=os.path.join(output_dir, "tsne_original_features.png"),
    )

    print("\nAnalysis complete.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Analyze microfacies separability in original well-log feature space."
    )
    parser.add_argument("--x_path", type=str, default="data/X.npy", help="Path to feature array (N, C)")
    parser.add_argument("--y_path", type=str, default="data/y.npy", help="Path to label array (N,)")
    parser.add_argument(
        "--output_dir",
        type=str,
        default="outputs/feature_separability",
        help="Directory for output figures",
    )
    parser.add_argument(
        "--max_samples",
        type=int,
        default=10000,
        help="Max samples for PCA/t-SNE visualization (default: 10000)",
    )
    parser.add_argument("--seed", type=int, default=42, help="Random seed")
    parser.add_argument(
        "--tsne_perplexity",
        type=float,
        default=30.0,
        help="t-SNE perplexity (auto-capped by sample size)",
    )

    # Optional xlsx export
    parser.add_argument(
        "--export_xlsx",
        type=str,
        default="",
        help="If set, export features/labels from this xlsx to --x_path/--y_path before analysis",
    )
    parser.add_argument(
        "--feature_cols",
        type=str,
        default="GR,CNL,DEN",
        help="Comma-separated feature columns for xlsx export",
    )
    parser.add_argument(
        "--label_col",
        type=str,
        default="facies",
        help="Label column name for xlsx export",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    if args.export_xlsx:
        feature_cols = [c.strip() for c in args.feature_cols.split(",") if c.strip()]
        export_xlsx_to_npy(
            xlsx_path=args.export_xlsx,
            x_path=args.x_path,
            y_path=args.y_path,
            feature_cols=feature_cols,
            label_col=args.label_col,
        )

    if not os.path.isfile(args.x_path) or not os.path.isfile(args.y_path):
        raise FileNotFoundError(
            f"Missing npy files: {args.x_path}, {args.y_path}\n"
            "Prepare them manually or pass --export_xlsx <path/to/data.xlsx>."
        )

    X, y = load_data(args.x_path, args.y_path)
    print(f"Loaded X {X.shape}, y {y.shape}")
    run_analysis(
        X=X,
        y=y,
        output_dir=args.output_dir,
        max_samples=args.max_samples,
        seed=args.seed,
        tsne_perplexity=args.tsne_perplexity,
    )


if __name__ == "__main__":
    main()
