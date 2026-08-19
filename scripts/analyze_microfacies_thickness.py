#!/usr/bin/env python3
"""统计 Excel 中 facies=0/1/2/3 连续层段的厚度并绘制频率直方图。

用法:
    python3 analyze_microfacies_thickness.py \
        --input facies-gr-diff0614-用GR-CNL-DEN.xlsx
"""
import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def runs_for_sheet(ws):
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    names = {str(v).strip(): i for i, v in enumerate(header) if v is not None}
    if "DEPT" not in names or "facies" not in names:
        return []
    di, fi = names["DEPT"], names["facies"]
    runs, current = [], None
    for row in rows:
        try:
            depth = float(row[di])
        except (TypeError, ValueError, IndexError):
            continue
        raw_facies = row[fi]
        try:
            facies_id = int(float(raw_facies))
        except (TypeError, ValueError):
            current = None
            continue
        if facies_id not in (0, 1, 2, 3):
            current = None
            continue
        if current is None or current["facies"] != facies_id:
            if current is not None:
                runs.append(current)
            current = {"well": ws.title, "facies": facies_id,
                       "top_depth": depth, "bottom_depth": depth,
                       "point_count": 1}
        else:
            current["bottom_depth"] = depth
            current["point_count"] += 1
    if current is not None:
        runs.append(current)
    for item in runs:
        item["thickness"] = abs(item["bottom_depth"] - item["top_depth"])
    return runs


def main():
    try:
        import matplotlib.pyplot as plt
        from matplotlib.font_manager import FontProperties
    except ImportError as exc:
        raise SystemExit("缺少绘图依赖 matplotlib，请先运行: pip install matplotlib") from exc
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--output-dir", type=Path, default=Path("microfacies_thickness"))
    ap.add_argument("--bins", type=int, default=20)
    args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(args.input, read_only=True, data_only=True)
    runs = [r for ws in wb.worksheets for r in runs_for_sheet(ws)]
    if not runs:
        raise SystemExit("没有找到同时包含 DEPT 和 facies 列的有效数据。")

    labels = {0: "河道", 1: "河口坝", 2: "席状砂", 3: "泥"}
    fields = ["well", "facies", "microfacies", "top_depth", "bottom_depth", "thickness", "point_count"]
    for r in runs:
        r["microfacies"] = labels[r["facies"]]
    with (args.output_dir / "microfacies_intervals.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(runs)
    grouped = defaultdict(list)
    for r in runs:
        grouped[r["microfacies"]].append(r["thickness"])
    with (args.output_dir / "microfacies_thickness_summary.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["microfacies", "interval_count", "total_thickness", "mean", "median", "min", "max"])
        for name, vals in sorted(grouped.items()):
            sv = sorted(vals); n = len(sv); med = sv[n//2] if n % 2 else (sv[n//2-1]+sv[n//2])/2
            w.writerow([name, n, sum(vals), sum(vals)/n, med, min(vals), max(vals)])

    # 直接指定字体文件，避免仅设置字体名称但系统未安装时中文仍乱码。
    font_candidates = [
        "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "C:/Windows/Fonts/msyh.ttc",
    ]
    font_path = next((p for p in font_candidates if Path(p).exists()), None)
    font_prop = FontProperties(fname=font_path) if font_path else FontProperties()
    if font_path:
        plt.rcParams["font.family"] = font_prop.get_name()
    plt.rcParams["axes.unicode_minus"] = False
    fig, axes = plt.subplots(2, 2, figsize=(20, 16), squeeze=False)
    colors = ["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728"]
    # 0--80 米范围内每 2 米一个箱，细分后更容易观察四类分布。
    for fid, color in zip(range(4), colors):
        ax = axes.flat[fid]
        values = grouped[labels[fid]]
        mean = sum(values) / len(values)
        variance = sum((v - mean) ** 2 for v in values) / len(values)
        ax.hist(values, bins=40, range=(0, 80), alpha=0.65,
                color=color, edgecolor="black")
        ax.set_xlabel("厚度区间（米）", fontproperties=font_prop, fontsize=28)
        ax.set_ylabel("频数", fontproperties=font_prop, fontsize=28)
        ax.set_title(f"{fid}: {labels[fid]}", fontproperties=font_prop, fontsize=30)
        ax.set_xlim(0, 50)
        ax.tick_params(axis="both", labelsize=28)
        legend_text = f"均值 = {mean:.2f} 米\n方差 = {variance:.2f} 米²"
        stat_prop = FontProperties(fname=font_path, size=22) if font_path else FontProperties(size=22)
        ax.legend([legend_text], loc="upper right", prop=stat_prop, framealpha=0.9)
    fig.suptitle("四种微相层段厚度频率直方图", fontproperties=font_prop, fontsize=32)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(args.output_dir / "facies_0_1_2_3_thickness_distribution.png", dpi=200); plt.close(fig)
    print(f"共统计 {len(runs)} 个连续微相层段、{len(grouped)} 种微相。结果目录: {args.output_dir}")


if __name__ == "__main__":
    main()
